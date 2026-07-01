#!/usr/bin/env python3
"""
Build the Kokoland 2024 reconciliation workbook (one .xlsx, multiple tabs)
from the consolidated JSON artifacts produced by build_summary.py + reconcile.py.
"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(__file__), "out")
def load(n): return json.load(open(os.path.join(OUT, n)))

HEAD = Font(bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F4E78")
WARN = PatternFill("solid", fgColor="FCE4D6")
BAD  = PatternFill("solid", fgColor="F8CBAD")
GOOD = PatternFill("solid", fgColor="E2EFDA")
BOLD = Font(bold=True)

def num(x): return x if isinstance(x, (int, float)) else 0.0

def sheet(wb, title, headers, rows, widths=None, fills=None):
    ws = wb.create_sheet(title)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h); cell.font = HEAD; cell.fill = HFILL
        cell.alignment = Alignment(horizontal="center")
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
        if fills:
            f = fills(row)
            if f:
                for c in range(1, len(headers) + 1):
                    ws.cell(r, c).fill = f
    ws.freeze_panes = "A2"
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = (widths or [16] * len(headers))[c - 1]
    return ws

def main():
    recs   = load("invoices_all.json")
    recon  = load("reconciliation.json")
    bank   = load("bank_ledger.json")
    adv    = load("advisor_figures.json")
    try:
        track = load("corrections_tracker.json")
    except FileNotFoundError:
        track = None

    # recompute buckets (mirror build_summary)
    BLOCKED = {"puneeth", "other_entity"}
    def purchase(r):
        return not (r.get("is_sales_invoice") or r.get("is_invoice") is False
                    or r.get("is_duplicate") or r.get("gross") is None)
    def bucket(r):
        s = r.get("recipient_status")
        if s in BLOCKED: return "BLOCKED"
        if num(r.get("gross")) <= 250: return "VALID"
        if s == "ok": return "VALID"
        return "AT_RISK"
    purchases = [r for r in recs if purchase(r)]
    for r in purchases: r["_b"] = bucket(r)

    valid   = sum(num(r["vat_amount"]) for r in purchases if r["_b"] == "VALID")
    atrisk  = sum(num(r["vat_amount"]) for r in purchases if r["_b"] == "AT_RISK")
    blocked = sum(num(r["vat_amount"]) for r in purchases if r["_b"] == "BLOCKED")
    total   = valid + atrisk + blocked

    wb = Workbook(); wb.remove(wb.active)

    # 1) Summary
    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 52; ws.column_dimensions["B"].width = 16
    ws["A1"] = "KOKOLAND GASTRO UG — 2024 Vorsteuer & Reconciliation"; ws["A1"].font = Font(bold=True, size=14)
    rows = [
        ("", ""),
        ("INPUT VAT (Vorsteuer)", "EUR"),
        ("  Valid now (named recipient or <=250 Kleinbetrag)", round(valid, 2)),
        ("  At risk (>250 & no named recipient — §14 gap)", round(atrisk, 2)),
        ("  Blocked (billed to Puneeth / other entity)", round(blocked, 2)),
        ("  TOTAL input VAT in documents", round(total, 2)),
        ("", ""),
        ("Recoverable with corrective action (at risk + blocked)", round(atrisk + blocked, 2)),
        ("", ""),
        ("DOCUMENTS", ""),
        ("  Purchase invoices counted", len(purchases)),
        ("  Excluded (sales/duplicate/non-invoice)", len(recs) - len(purchases)),
        ("", ""),
        ("BANK RECONCILIATION", ""),
        ("  Matched invoice <-> payment", len(recon["matched"])),
        ("  Unmatched invoices (no bank hit)", len(recon["unmatched_invoices"])),
        ("  Candidate MISSING invoices (bank debit, no doc)", len(recon["candidate_missing_invoices"])),
        ("  Internal/treasury debits excluded", recon.get("internal_debits_excluded", 0)),
    ]
    for i, (a, b) in enumerate(rows, 3):
        ws.cell(i, 1, a); ws.cell(i, 2, b)
        if isinstance(a, str) and a.isupper(): ws.cell(i, 1).font = BOLD
    ws.cell(8, 1).font = BOLD; ws.cell(8, 2).font = BOLD  # TOTAL
    ws.cell(10, 1).font = BOLD; ws.cell(10, 2).font = BOLD

    # 2) Invoices (consolidated)
    ihead = ["file", "week", "date", "vendor", "invoice_no", "net", "vat_amount",
             "vat_rate", "gross", "recipient_status", "bucket", "paid_cash", "category", "notes"]
    irows = []
    for r in recs:
        b = r["_b"] if r in purchases else ("EXCLUDED" if not purchase(r) else "")
        irows.append([r.get("file"), r.get("week"), r.get("date"), r.get("vendor"),
                      r.get("invoice_no"), r.get("net"), r.get("vat_amount"),
                      r.get("vat_rate"), r.get("gross"), r.get("recipient_status"),
                      b, r.get("paid_cash"), r.get("category"), r.get("notes")])
    def ifill(row):
        return {"BLOCKED": BAD, "AT_RISK": WARN, "VALID": GOOD}.get(row[10])
    sheet(wb, "Invoices", ihead, irows,
          widths=[34,6,11,26,14,9,9,6,9,14,9,8,18,60], fills=ifill)

    # 3) Wrong-name / blocked
    bhead = ["date", "vendor", "invoice_no", "net", "vat_amount", "gross", "recipient_status", "file", "notes"]
    brows = [[r["date"], r["vendor"], r.get("invoice_no"), r["net"], r["vat_amount"], r["gross"],
              r["recipient_status"], r["file"], r.get("notes")]
             for r in sorted(purchases, key=lambda x: -num(x["vat_amount"])) if r["_b"] == "BLOCKED"]
    sheet(wb, "Wrong-name (blocked VAT)", bhead, brows, widths=[11,26,14,9,9,9,14,32,60],
          fills=lambda r: BAD)

    # 4) §14 compliance gaps
    grows = [[r["date"], r["vendor"], r["gross"], r["vat_amount"], r["recipient_status"], r["file"]]
             for r in sorted(purchases, key=lambda x: -num(x["gross"])) if r["_b"] == "AT_RISK"]
    sheet(wb, "§14 gaps (>250 no name)", ["date","vendor","gross","vat_amount","recipient_status","file"],
          grows, widths=[11,30,10,10,14,34], fills=lambda r: WARN)

    # 5) Matched
    mrows = [[m["invoice"], m["gross"], m["bank_date"], m["bank_amount"], m["account"]]
             for m in recon["matched"]]
    sheet(wb, "Bank matched", ["invoice","gross","bank_date","bank_amount","account"], mrows,
          widths=[36,10,12,12,22])

    # 6) Unmatched invoices
    urows = [[u["date"], u["vendor"], u["gross"], u["file"]] for u in
             sorted(recon["unmatched_invoices"], key=lambda x: -num(x["gross"]))]
    sheet(wb, "Bank unmatched invoices", ["date","vendor","gross","file"], urows,
          widths=[11,30,10,40])

    # 7) Candidate missing invoices
    crows = [[c["date"], abs(num(c["amount"])), c["account"], c.get("counterparty"), c.get("description")]
             for c in sorted(recon["candidate_missing_invoices"], key=lambda x: num(x["amount"]))]
    sheet(wb, "Candidate MISSING invoices", ["date","amount","account","counterparty","description"],
          crows, widths=[11,11,22,30,50])

    # 7b) Related-party transfers
    rp = recon.get("related_party_transfers", [])
    rprows = [[r["date"], abs(num(r["amount"])), r["account"], r.get("counterparty"), r.get("description")]
              for r in sorted(rp, key=lambda x: num(x["amount"]))]
    sheet(wb, "Related-party transfers", ["date","amount","account","counterparty","description"],
          rprows, widths=[11,11,22,30,50], fills=lambda r: WARN)

    # 8) Bank ledger (full)
    lhead = ["date","account","type","amount","balance","counterparty","description","category"]
    lrows = [[b["date"], b["account"], b["type"], b["amount"], b.get("balance"),
              b.get("counterparty"), b.get("description"), b.get("category")] for b in bank]
    sheet(wb, "Bank ledger (full)", lhead, lrows, widths=[11,20,10,11,11,28,46,16])

    # 9) Advisor verification
    adv_vat = adv["ust_erklaerung"]["input_vat_total"]
    gap = round(adv_vat - total, 2)
    ws = wb.create_sheet("Advisor verification")
    ws.column_dimensions["A"].width = 50
    for c in (2, 3): ws.column_dimensions[get_column_letter(c)].width = 16
    ws["A1"] = "ADVISOR (Hegazi & Kollegen) vs THIS ENGINE"; ws["A1"].font = Font(bold=True, size=13)
    vrows = [
        ("Metric", "Advisor filed", "This engine"),
        ("Input VAT / Vorsteuer (EUR)", adv_vat, round(total, 2)),
        ("  — substantiated by documents on hand", "", round(total, 2)),
        ("  — of which §14-clean (claimable)", "", round(valid, 2)),
        ("  — of which billed to Puneeth/other entity", "", round(blocked, 2)),
        ("  — of which §14 gap (>250 no name)", "", round(atrisk, 2)),
        ("UNRECONCILED VAT GAP (advisor − engine)", gap, ""),
        ("", "", ""),
        ("Output VAT (EUR)", adv["ust_erklaerung"]["output_vat_total"], ""),
        ("VAT refund filed (EUR)", adv["ust_erklaerung"]["refund"], ""),
        ("", "", ""),
        ("Revenue / Umsatzerlöse", adv["guv"]["umsatzerloese"], ""),
        ("Renovation (Instandhaltung 4260)", adv["guv"]["instandhaltung_raeume_4260"], ""),
        ("Net loss (Jahresfehlbetrag)", adv["guv"]["jahresfehlbetrag"], ""),
        ("", "", ""),
        ("Advisor fee — UG (Rech 145/2026)", adv["advisor_fees"]["rechnung_145_2026_UG"], ""),
        ("Advisor fee — personal ESt (Rech 1894/2025)", adv["advisor_fees"]["rechnung_1894_2025_personal_ESt"], ""),
        ("", "", ""),
        ("INTERPRETATION", "", ""),
        ("The €%.0f gap = renovation/subcontractor spend booked by" % gap, "", ""),
        ("the advisor (e.g. Malerbetrieb Lampert, plumbers) for which", "", ""),
        ("no invoice is in the provided weekly folders. Locate those", "", ""),
        ("invoices to substantiate the refund; re-bill Puneeth invoices", "", ""),
        ("to Kokoland to protect €%.2f of input VAT." % blocked, "", ""),
    ]
    for i, (a, b, c) in enumerate(vrows, 3):
        ws.cell(i, 1, a); ws.cell(i, 2, b); ws.cell(i, 3, c)
        if a in ("Metric", "INTERPRETATION") or "GAP" in str(a):
            for cc in (1, 2, 3): ws.cell(i, cc).font = BOLD
    ws.cell(9, 1).fill = WARN; ws.cell(9, 2).fill = WARN

    # 10) Corrections tracker
    if track:
        GREEN = PatternFill("solid", fgColor="C6EFCE")
        YELLOW = PatternFill("solid", fgColor="FFEB9C")
        thead = ["id", "stream", "vendor", "VAT at stake", "email sent", "corrected received", "status", "notes"]
        trows = []
        for w in track["workstreams"]:
            trows.append([w["id"], w["stream"], w["vendor"], w.get("vat_at_stake", 0),
                          w.get("email_sent"), w.get("corrected_received"), w["status"], w.get("notes")])
        def tfill(row):
            return {"DONE": GREEN, "PARTIAL": YELLOW, "EMAIL_SENT": YELLOW, "TODO": BAD}.get(row[6])
        ws = sheet(wb, "Corrections tracker", thead, trows,
                   widths=[5, 22, 34, 12, 11, 16, 9, 60], fills=tfill)
        done = sum(w.get("vat_at_stake", 0) for w in track["workstreams"] if w["status"] == "DONE")
        openv = sum(w.get("vat_at_stake", 0) for w in track["workstreams"] if w["status"] in ("TODO", "PARTIAL", "EMAIL_SENT"))
        r = len(trows) + 3
        ws.cell(r, 3, "VAT secured (DONE):"); ws.cell(r, 4, round(done, 2)); ws.cell(r, 3).font = BOLD
        ws.cell(r + 1, 3, "VAT still open:"); ws.cell(r + 1, 4, round(openv, 2)); ws.cell(r + 1, 3).font = BOLD

    path = os.path.join(OUT, "Kokoland_2024_Reconciliation.xlsx")
    wb.save(path)
    print("wrote", path)
    print("tabs:", wb.sheetnames)

if __name__ == "__main__":
    main()
